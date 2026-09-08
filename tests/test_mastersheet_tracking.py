from __future__ import annotations

import dataclasses
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from agent.config import load_settings, load_story_bank
from agent.db import connect, init_db
from agent.mastersheet_tracking import (
    CANDIDATES_HEADER,
    CANDIDATES_SHEET,
    count_new_candidates,
    import_next_batch,
    read_raw_candidates,
    sync_pipeline_candidates,
)


def _build_mastersheet(path: Path, raw_rows: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "raw"
    ws.append(["Company", "Name", "Designation", "E-Mail 1", "Email 2", "Notes"])
    for row in raw_rows:
        ws.append(row)
    wb.save(path)


def _build_candidates_sheet(path: Path, rows: list[tuple]) -> None:
    """rows are (Company, Contact Name, Designation, Email, Domain, Status, First Seen, Note)."""
    wb = openpyxl.Workbook()
    raw_ws = wb.active
    raw_ws.title = "raw"  # import_next_batch doesn't read 'raw' -- just needs the file valid
    ws = wb.create_sheet(CANDIDATES_SHEET)
    ws.append(CANDIDATES_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(path)


@pytest.fixture
def mastersheet_path(tmp_path: Path) -> Path:
    return tmp_path / "Email Mastersheet.xlsx"


@pytest.fixture
def settings(tmp_path: Path):
    real_settings = load_settings()
    return dataclasses.replace(real_settings, db_path=tmp_path / "agent.db")


@pytest.fixture
def story_bank():
    return load_story_bank()


def test_read_raw_candidates_skips_rows_with_no_company_or_no_email(mastersheet_path):
    _build_mastersheet(
        mastersheet_path,
        [
            ("Acme", "Jane Doe", "PM", "jane@acme.com", None, "a note"),
            ("", "No Company", "PM", "noco@example.com", None, None),  # no company
            ("NoEmail Inc", "No Email", "PM", None, None, None),  # no email at all
            ("Fallback Co", "Uses Email2", "PM", None, "fallback@fallback.com", None),
        ],
    )
    wb = openpyxl.load_workbook(mastersheet_path)

    candidates = read_raw_candidates(wb)

    emails = {c.email for c in candidates}
    assert emails == {"jane@acme.com", "fallback@fallback.com"}


def test_sync_pipeline_candidates_creates_sheet_and_adds_new_rows(mastersheet_path):
    _build_mastersheet(
        mastersheet_path,
        [("Acme", "Jane Doe", "PM", "jane@acme.com", None, None)],
    )

    stats = sync_pipeline_candidates(
        mastersheet_path, existing_domains=set(), existing_names=set(), today=date(2026, 9, 9)
    )

    assert stats.candidates_found == 1
    assert stats.newly_tracked == 1
    assert stats.already_in_pipeline_db == 0
    assert stats.already_tracked == 0

    wb = openpyxl.load_workbook(mastersheet_path)
    assert CANDIDATES_SHEET in wb.sheetnames
    rows = list(wb[CANDIDATES_SHEET].iter_rows(min_row=2, values_only=True))
    assert rows == [("Acme", "Jane Doe", "PM", "jane@acme.com", "acme.com", "New", "2026-09-09", None)]


def test_sync_pipeline_candidates_skips_company_already_in_agent_db_by_domain(mastersheet_path):
    _build_mastersheet(
        mastersheet_path,
        [("Acme", "Jane Doe", "PM", "jane@acme.com", None, None)],
    )

    stats = sync_pipeline_candidates(
        mastersheet_path, existing_domains={"acme.com"}, existing_names=set()
    )

    assert stats.already_in_pipeline_db == 1
    assert stats.newly_tracked == 0


def test_sync_pipeline_candidates_skips_company_already_in_agent_db_by_name_fallback(mastersheet_path):
    """Some agent.db companies have no domain (e.g. hand-entered test rows) -- name
    match is the fallback dedup path for those."""
    _build_mastersheet(
        mastersheet_path,
        [("No Domain Co", "Jane Doe", "PM", "jane@nodomainco.com", None, None)],
    )

    stats = sync_pipeline_candidates(
        mastersheet_path, existing_domains=set(), existing_names={"no domain co"}
    )

    assert stats.already_in_pipeline_db == 1
    assert stats.newly_tracked == 0


def test_sync_pipeline_candidates_is_idempotent_across_reruns(mastersheet_path):
    _build_mastersheet(
        mastersheet_path,
        [("Acme", "Jane Doe", "PM", "jane@acme.com", None, None)],
    )

    first = sync_pipeline_candidates(mastersheet_path, existing_domains=set(), existing_names=set())
    second = sync_pipeline_candidates(mastersheet_path, existing_domains=set(), existing_names=set())

    assert first.newly_tracked == 1
    assert second.newly_tracked == 0
    assert second.already_tracked == 1

    wb = openpyxl.load_workbook(mastersheet_path)
    rows = list(wb[CANDIDATES_SHEET].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1  # not duplicated


def test_sync_pipeline_candidates_leaves_other_sheets_untouched(mastersheet_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "raw"
    ws.append(["Company", "Name", "Designation", "E-Mail 1", "Email 2", "Notes"])
    ws.append(("Acme", "Jane Doe", "PM", "jane@acme.com", None, None))
    company_sheet = wb.create_sheet("Company")
    company_sheet.append(["Company Name", None, "Status", "Comments"])
    company_sheet.append(["Air India", None, "Did Not Reply", None])
    wb.save(mastersheet_path)

    sync_pipeline_candidates(mastersheet_path, existing_domains=set(), existing_names=set())

    wb = openpyxl.load_workbook(mastersheet_path)
    company_rows = list(wb["Company"].iter_rows(min_row=2, values_only=True))
    assert company_rows == [("Air India", None, "Did Not Reply", None)]


# ---------------------------------------------------------------------------
# count_new_candidates
# ---------------------------------------------------------------------------


def test_count_new_candidates_counts_only_status_new(mastersheet_path):
    _build_candidates_sheet(
        mastersheet_path,
        [
            ("Acme", "Jane", "PM", "jane@acme.com", "acme.com", "New", "2026-09-08", None),
            ("Beta", "Bob", "PM", "bob@beta.com", "beta.com", "Imported", "2026-09-08", None),
            ("Gamma", "Gia", "PM", "gia@gamma.com", "gamma.com", "New", "2026-09-08", None),
        ],
    )

    assert count_new_candidates(mastersheet_path) == 2


def test_count_new_candidates_missing_file_returns_zero(tmp_path):
    assert count_new_candidates(tmp_path / "does_not_exist.xlsx") == 0


def test_count_new_candidates_missing_sheet_returns_zero(mastersheet_path):
    _build_mastersheet(mastersheet_path, [])  # only a 'raw' sheet, no Pipeline Candidates
    assert count_new_candidates(mastersheet_path) == 0


# ---------------------------------------------------------------------------
# import_next_batch
# ---------------------------------------------------------------------------


def test_import_next_batch_creates_company_contact_draft_and_marks_imported(
    mastersheet_path, settings, story_bank
):
    _build_candidates_sheet(
        mastersheet_path,
        [
            (
                "TestPay", "Jane Doe", "Product Manager", "jane@testpay.com", "testpay.com",
                "New", "2026-09-08", "great UPI payments experience",
            ),
        ],
    )
    init_db(settings.db_path, settings)

    with connect(settings.db_path) as conn:
        stats = import_next_batch(
            conn,
            mastersheet_path,
            story_bank=story_bank,
            templates_dir=settings.templates_dir,
            sender_display_name=settings.sender_display_name,
            attach_resume_by_default=settings.attach_resume_by_default,
            sender_phone=settings.sender_phone,
            niche_keywords={"fintech": ["payments", "upi"], "consumer": []},
            niche_order=["consumer", "fintech"],
        )

        company = conn.execute("SELECT * FROM companies WHERE name = 'TestPay'").fetchone()
        contact = conn.execute("SELECT * FROM contacts WHERE email = 'jane@testpay.com'").fetchone()
        draft = conn.execute(
            "SELECT * FROM email_queue WHERE contact_id = ?", (contact["id"],)
        ).fetchone()

    assert stats.candidates_pulled == 1
    assert stats.companies_created == 1
    assert stats.contacts_inserted == 1
    assert stats.companies_classified == 1
    assert stats.drafts_generated == 1
    assert stats.remaining_new == 0

    assert company["niche"] == "fintech"
    assert contact["status"] == "verified"
    assert contact["role_category"] == "product"
    assert draft is not None
    assert draft["status"] == "pending_review"

    wb = openpyxl.load_workbook(mastersheet_path)
    row = next(wb[CANDIDATES_SHEET].iter_rows(min_row=2, values_only=True))
    assert row[5] == "Imported"  # Status column


def test_import_next_batch_respects_batch_size(mastersheet_path, settings, story_bank):
    rows = [
        (f"Co{i}", f"Person{i}", "PM", f"p{i}@co{i}.com", f"co{i}.com", "New", "2026-09-08", None)
        for i in range(15)
    ]
    _build_candidates_sheet(mastersheet_path, rows)
    init_db(settings.db_path, settings)

    with connect(settings.db_path) as conn:
        stats = import_next_batch(
            conn,
            mastersheet_path,
            story_bank=story_bank,
            templates_dir=settings.templates_dir,
            sender_display_name=settings.sender_display_name,
            attach_resume_by_default=settings.attach_resume_by_default,
            sender_phone=settings.sender_phone,
            niche_keywords={},
            niche_order=[],
            batch_size=10,
        )

    assert stats.candidates_pulled == 10
    assert stats.remaining_new == 5
    assert count_new_candidates(mastersheet_path) == 5


def test_import_next_batch_skips_contact_already_in_agent_db(mastersheet_path, settings, story_bank):
    _build_candidates_sheet(
        mastersheet_path,
        [("Acme", "Jane", "PM", "jane@acme.com", "acme.com", "New", "2026-09-08", None)],
    )
    init_db(settings.db_path, settings)

    with connect(settings.db_path) as conn:
        company_id = conn.execute(
            "INSERT INTO companies (name, domain, source) VALUES ('Acme', 'acme.com', 'manual_csv')"
        ).lastrowid
        conn.execute(
            "INSERT INTO contacts (company_id, email, status) VALUES (?, 'jane@acme.com', 'verified')",
            (company_id,),
        )
        conn.commit()

        stats = import_next_batch(
            conn,
            mastersheet_path,
            story_bank=story_bank,
            templates_dir=settings.templates_dir,
            sender_display_name=settings.sender_display_name,
            attach_resume_by_default=settings.attach_resume_by_default,
            sender_phone=settings.sender_phone,
            niche_keywords={},
            niche_order=[],
        )

    assert stats.candidates_pulled == 1
    assert stats.companies_matched == 1  # matched the existing Acme by domain
    assert stats.contacts_inserted == 0
    assert stats.contacts_skipped_duplicate == 1


def test_import_next_batch_no_candidates_sheet_returns_empty_stats(mastersheet_path, settings, story_bank):
    _build_mastersheet(mastersheet_path, [])  # only 'raw', no Pipeline Candidates sheet
    init_db(settings.db_path, settings)

    with connect(settings.db_path) as conn:
        stats = import_next_batch(
            conn,
            mastersheet_path,
            story_bank=story_bank,
            templates_dir=settings.templates_dir,
            sender_display_name=settings.sender_display_name,
            attach_resume_by_default=settings.attach_resume_by_default,
            sender_phone=settings.sender_phone,
            niche_keywords={},
            niche_order=[],
        )

    assert stats.candidates_pulled == 0
