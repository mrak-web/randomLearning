from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from agent.linkedin_posts import (
    ApifyLinkedInPostSearch,
    HiringPost,
    LinkedInPostSearchClient,
    LinkedInPostSearchError,
    author_matches_role_signal,
    build_post_search_queries,
    filter_relevant_posts,
    has_experience_gap,
    has_seniority_mismatch,
    looks_like_hiring_post,
    populate_hiring_posts_sheet,
    search_hiring_posts,
    write_hiring_posts_excel,
)

TITLE_KEYWORDS = ["Product Manager", "APM"]

# ---------------------------------------------------------------------------
# Fakes standing in for the real HTTP layer / a real LinkedInPostSearchClient
# ---------------------------------------------------------------------------


class FakeResponse:
    def __init__(self, status_code: int = 200, json_data=None):
        self.status_code = status_code
        self._json_data = json_data if json_data is not None else []

    def json(self):
        return self._json_data


class FakeSession:
    def __init__(self, response_by_query: dict):
        self._response_by_query = response_by_query
        self.calls: list[dict] = []

    def post(self, url, params=None, json=None, timeout=None):
        self.calls.append({"url": url, "params": params, "json": json, "timeout": timeout})
        query = json["searchQueries"][0]
        result = self._response_by_query[query]
        if isinstance(result, Exception):
            raise result
        return result


def raw_item(
    post_id: str,
    content: str = "We're hiring a Product Manager, DM for referral.",
    author_name: str = "Jane Doe",
    author_info: str = "Product Manager @ Acme",
    link: str | None = None,
    company_name: str | None = "Acme",
) -> dict:
    attributes = [{"company": {"name": company_name}}] if company_name else []
    return {
        "id": post_id,
        "content": content,
        "author": {
            "name": author_name,
            "info": author_info,
            "linkedinUrl": f"https://in.linkedin.com/in/{author_name.replace(' ', '-').lower()}",
        },
        "linkedinUrl": link or f"https://www.linkedin.com/posts/{post_id}",
        "postedAt": {"date": "2026-08-20"},
        "contentAttributes": attributes,
    }


class FakePostSearchClient(LinkedInPostSearchClient):
    def __init__(self, responses: dict):
        self.responses = responses
        self.calls: list[tuple] = []

    def search(self, query, posted_limit, limit):
        self.calls.append((query, posted_limit, limit))
        result = self.responses.get(query, [])
        if isinstance(result, Exception):
            raise result
        return result[:limit]


def hiring_post(
    post_id: str,
    link: str | None = None,
    content: str = "We're hiring a Product Manager, DM for referral.",
    author_headline: str = "Product Manager @ Acme",
) -> HiringPost:
    return HiringPost(
        post_id=post_id,
        author_name="Jane Doe",
        author_headline=author_headline,
        author_profile_url="https://in.linkedin.com/in/jane-doe",
        content=content,
        link=link or f"https://www.linkedin.com/posts/{post_id}",
        posted_at="2026-08-20",
        company_name="Acme",
    )


# ---------------------------------------------------------------------------
# build_post_search_queries
# ---------------------------------------------------------------------------


def test_build_post_search_queries_default_locations_includes_location_free_variant():
    queries = build_post_search_queries(["Product Manager", "APM"], locations=["India", "Bengaluru"])

    # Location-free variant for every title comes first (location-slot-major order).
    assert queries[:2] == ['"Product Manager" hiring', '"APM" hiring']
    assert '"Product Manager" hiring India' in queries
    assert '"APM" hiring India' in queries
    assert '"Product Manager" hiring Bengaluru' in queries
    assert '"APM" hiring Bengaluru' in queries
    assert len(queries) == 2 * 3  # 2 titles x (1 location-free + 2 locations)


def test_build_post_search_queries_uses_default_locations_when_omitted():
    queries = build_post_search_queries(["Product Manager"])

    assert '"Product Manager" hiring' in queries
    assert '"Product Manager" hiring India' in queries
    assert '"Product Manager" hiring Bengaluru' in queries
    assert '"Product Manager" hiring Bangalore' in queries


# ---------------------------------------------------------------------------
# looks_like_hiring_post / author_matches_role_signal / filter_relevant_posts
# ---------------------------------------------------------------------------


def test_looks_like_hiring_post_requires_both_intent_and_title():
    assert looks_like_hiring_post(
        "We're hiring a Product Manager for our team", TITLE_KEYWORDS
    )
    # Hiring intent but no matching title keyword.
    assert not looks_like_hiring_post("We're hiring a Sales Executive", TITLE_KEYWORDS)
    # Title keyword present but no hiring-intent phrase (e.g. a candidate's own bio).
    assert not looks_like_hiring_post(
        "I am a Product Manager with 5 years experience", TITLE_KEYWORDS
    )
    assert not looks_like_hiring_post(None, TITLE_KEYWORDS)


def test_looks_like_hiring_post_matches_real_examples_shared_by_arjun():
    # post1 (2026-08-23): first-person hiring post naming cities, not "India".
    post1 = (
        "I am hiring a Product Manager for roles in London and Bangalore.\n\n"
        "We're looking for someone with experience building digital, user-facing "
        "products who understands customer needs and can turn ideas into simple, "
        "useful experiences."
    )
    assert looks_like_hiring_post(post1, TITLE_KEYWORDS)

    # post3 (2026-08-23): explicit "Associate Product Manager" hiring post.
    post3 = (
        "I am hiring an Associate Product Manager at Saber.\n\n"
        "Saber is a YC-backed payments company moving millions of dollars in "
        "cross-border remittance every day."
    )
    assert looks_like_hiring_post(post3, TITLE_KEYWORDS)

    # post2 (2026-08-23, Kirana Club): accepted miss -- never says "Product
    # Manager" literally, only lists "Product" as one category among several
    # ("Product, Growth, Engineering and Business"). Loosening the title match to
    # bare "Product" would catch far too many unrelated posts, so this stays a
    # known gap for the keyword heuristic rather than reached-for with NLP.
    post2 = (
        "Kirana Club is entering its next phase of growth — and we're looking "
        "for builders.\nWe have multiple roles open across Product, Growth, "
        "Engineering and Business."
    )
    assert not looks_like_hiring_post(post2, TITLE_KEYWORDS)


def test_looks_like_hiring_post_ignores_incidental_mention_deep_in_body():
    # Real example from manual testing: a recruiter's post hiring a Frontend
    # Developer that happens to mention "product managers" as a collaborating role
    # deep in the body -- should NOT count as a Product Manager hiring post.
    content = (
        "We're Hiring! Frontend Developer (React & Modern JavaScript)\n"
        + ("Filler text to push past the headline window. " * 10)
        + "You'll work directly with UI/UX designers, backend engineers, and "
        "product managers to turn wireframes into fast, accessible features."
    )
    assert not looks_like_hiring_post(content, TITLE_KEYWORDS)


def test_author_matches_role_signal():
    assert author_matches_role_signal("Product Manager @ Acme")
    assert author_matches_role_signal("Talent Acquisition Specialist")
    assert author_matches_role_signal("Founder & CEO")
    assert not author_matches_role_signal("159 followers")
    assert not author_matches_role_signal(None)


SENIORITY_MISMATCH_KEYWORDS = ["senior", "principal", "director"]


def test_has_seniority_mismatch():
    assert has_seniority_mismatch(
        "Product Manager – Enterprise Wi-Fi\nExperience: 8-12 Years\nSenior candidates preferred",
        SENIORITY_MISMATCH_KEYWORDS,
    )
    assert not has_seniority_mismatch("Associate Product Manager, 1-2 years", SENIORITY_MISMATCH_KEYWORDS)
    assert not has_seniority_mismatch(None, SENIORITY_MISMATCH_KEYWORDS)
    assert not has_seniority_mismatch("Product Manager role", [])


def test_has_experience_gap():
    # Real example (2026-08-23): no seniority word, only a numeric years bar.
    content = "Product Manager - Enterprise Wi-Fi\nExperience: 8-12 Years"
    assert has_experience_gap(content, total_experience_years=2.2)
    assert not has_experience_gap(content, total_experience_years=9.0)
    # Small gap (<= default max_gap_years) is left for a human to judge, not dropped.
    assert not has_experience_gap("Requires 3+ years", total_experience_years=2.2)
    assert not has_experience_gap(None, total_experience_years=2.2)
    assert not has_experience_gap("No years mentioned here", total_experience_years=2.2)


def test_filter_relevant_posts_requires_all_conditions():
    posts = [
        hiring_post("relevant"),  # hiring intent + title + product headline -> keep
        hiring_post(
            "wrong_role", content="We're hiring a Sales Executive"
        ),  # no title match -> drop
        hiring_post(
            "bot_account", author_headline="159 followers"
        ),  # no author role signal -> drop
        hiring_post(
            "candidate_bio",
            content="I am a Product Manager looking for my next role",
            author_headline="Product Manager @ Acme",
        ),  # no hiring intent -> drop
        hiring_post(
            "senior_role",
            content="We're hiring a Senior Product Manager, 8-12 years experience",
        ),  # seniority word mismatch -> drop
        hiring_post(
            "years_gap",
            content="Hiring a Product Manager. Experience: 8-12 Years required.",
        ),  # no seniority word, but years bar far exceeds Arjun's -> drop
    ]

    filtered = filter_relevant_posts(
        posts, TITLE_KEYWORDS, SENIORITY_MISMATCH_KEYWORDS, total_experience_years=2.2
    )

    assert [p.post_id for p in filtered] == ["relevant"]


def test_filter_relevant_posts_seniority_and_experience_checks_default_to_off():
    # A caller that doesn't pass the new params keeps prior behavior -- a senior
    # role/large years bar still passes if it clears the other two conditions.
    posts = [
        hiring_post("senior_role", content="We're hiring a Senior Product Manager"),
        hiring_post("years_gap", content="Hiring a Product Manager, 8-12 Years required"),
    ]

    filtered = filter_relevant_posts(posts, TITLE_KEYWORDS)

    assert [p.post_id for p in filtered] == ["senior_role", "years_gap"]


# ---------------------------------------------------------------------------
# ApifyLinkedInPostSearch
# ---------------------------------------------------------------------------


def test_apify_post_search_requires_api_token():
    with pytest.raises(LinkedInPostSearchError, match="token"):
        ApifyLinkedInPostSearch(api_token="")


def test_apify_post_search_parses_items():
    session = FakeSession(
        {
            '"Product Manager" hiring India': FakeResponse(
                json_data=[raw_item("post1"), raw_item("post2", company_name=None)]
            )
        }
    )
    client = ApifyLinkedInPostSearch(api_token="tok", session=session)

    results = client.search('"Product Manager" hiring India', "month", 10)

    assert len(results) == 2
    assert results[0].author_name == "Jane Doe"
    assert results[0].company_name == "Acme"
    assert results[1].company_name is None

    call = session.calls[0]
    assert call["params"] == {"token": "tok"}
    assert call["json"]["searchQueries"] == ['"Product Manager" hiring India']
    assert call["json"]["postedLimit"] == "month"
    assert call["json"]["maxPosts"] == 10


def test_apify_post_search_zero_limit_short_circuits():
    session = FakeSession({})
    client = ApifyLinkedInPostSearch(api_token="tok", session=session)

    assert client.search('"Product Manager" hiring India', "month", 0) == []
    assert session.calls == []


def test_apify_post_search_non_200_raises():
    session = FakeSession({'"Product Manager" hiring India': FakeResponse(status_code=500)})
    client = ApifyLinkedInPostSearch(api_token="tok", session=session)

    with pytest.raises(LinkedInPostSearchError, match="HTTP 500"):
        client.search('"Product Manager" hiring India', "month", 10)


# ---------------------------------------------------------------------------
# search_hiring_posts
# ---------------------------------------------------------------------------


def test_search_hiring_posts_caps_combined_total():
    client = FakePostSearchClient(
        {
            "q1": [hiring_post("p1"), hiring_post("p2"), hiring_post("p3")],
            "q2": [hiring_post("p4"), hiring_post("p5")],
        }
    )

    results = search_hiring_posts(client, queries=["q1", "q2"], posted_limit="month", total_limit=4)

    assert [p.post_id for p in results] == ["p1", "p2", "p3", "p4"]
    assert client.calls == [("q1", "month", 4), ("q2", "month", 1)]


def test_search_hiring_posts_dedupes_by_link():
    shared = hiring_post("shared", link="https://www.linkedin.com/posts/shared")
    client = FakePostSearchClient(
        {
            "q1": [shared],
            "q2": [shared, hiring_post("unique")],
        }
    )

    results = search_hiring_posts(client, queries=["q1", "q2"], posted_limit="month", total_limit=10)

    assert [p.post_id for p in results] == ["shared", "unique"]


# ---------------------------------------------------------------------------
# write_hiring_posts_excel / populate_hiring_posts_sheet
# ---------------------------------------------------------------------------


def test_write_hiring_posts_excel_round_trips(tmp_path: Path):
    posts = [
        HiringPost(
            post_id="post1",
            author_name="Jane Doe",
            author_headline="Product Manager @ Acme",
            author_profile_url="https://in.linkedin.com/in/jane-doe",
            content="We're hiring a Product Manager, DM for referral." * 20,
            link="https://www.linkedin.com/posts/post1",
            posted_at="2026-08-20",
            company_name="Acme",
        )
    ]

    output_path = tmp_path / "posts.xlsx"
    write_hiring_posts_excel(posts, output_path)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.title == "Hiring Posts"

    header = [cell.value for cell in sheet[1]]
    assert header == [
        "Contact Name",
        "Contact Headline",
        "Contact LinkedIn Profile",
        "Company (if detected)",
        "Post Excerpt",
        "Post Link",
        "Posted At",
    ]

    row = [cell.value for cell in sheet[2]]
    assert row[0] == "Jane Doe"
    assert row[3] == "Acme"
    assert row[4].endswith("…")  # long content gets truncated with an ellipsis
    assert len(row[4]) <= 601

    assert sheet.cell(row=2, column=3).hyperlink.target == "https://in.linkedin.com/in/jane-doe"
    assert sheet.cell(row=2, column=6).hyperlink.target == "https://www.linkedin.com/posts/post1"


def test_write_hiring_posts_excel_empty_list_still_writes_header(tmp_path: Path):
    output_path = tmp_path / "posts.xlsx"
    write_hiring_posts_excel([], output_path)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.max_row == 1
