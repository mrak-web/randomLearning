from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from agent.linkedin_jobs import (
    ApifyLinkedInJobSearch,
    LinkedInJobPosting,
    LinkedInJobSearchClient,
    LinkedInJobSearchError,
    extract_emails,
    search_jobs_for_keywords,
    write_postings_excel,
)

# ---------------------------------------------------------------------------
# Fakes standing in for the real HTTP layer / a real LinkedInJobSearchClient
# ---------------------------------------------------------------------------


class FakeResponse:
    def __init__(self, status_code: int = 200, json_data=None):
        self.status_code = status_code
        self._json_data = json_data if json_data is not None else []

    def json(self):
        return self._json_data


class FakeSession:
    """Stands in for requests.Session; records calls, returns/raises canned results."""

    def __init__(self, response_by_keyword: dict):
        self._response_by_keyword = response_by_keyword
        self.calls: list[dict] = []

    def post(self, url, params=None, json=None, timeout=None):
        self.calls.append({"url": url, "params": params, "json": json, "timeout": timeout})
        keyword = json["keywords"]
        result = self._response_by_keyword[keyword]
        if isinstance(result, Exception):
            raise result
        return result


def raw_item(
    job_id: str,
    title: str = "Product Manager",
    company_name: str = "Acme",
    link: str | None = None,
    poster_name: str | None = None,
    description: str | None = None,
) -> dict:
    return {
        "id": job_id,
        "title": title,
        "companyName": company_name,
        "link": link or f"https://in.linkedin.com/jobs/view/{job_id}",
        "location": "Bengaluru, India",
        "postedAt": "2026-08-20",
        "applicantsCount": "50",
        "jobPosterName": poster_name,
        "jobPosterTitle": "Senior HR Executive" if poster_name else None,
        "jobPosterProfileUrl": f"https://in.linkedin.com/in/{poster_name}" if poster_name else None,
        "descriptionText": description,
    }


class FakeJobSearchClient(LinkedInJobSearchClient):
    def __init__(self, responses: dict):
        self.responses = responses
        self.calls: list[tuple] = []

    def search(self, keywords, location, date_posted, limit):
        self.calls.append((keywords, location, date_posted, limit))
        result = self.responses.get(keywords, [])
        if isinstance(result, Exception):
            raise result
        return result[:limit]


def posting(job_id: str, link: str | None = None, poster_name: str | None = None) -> LinkedInJobPosting:
    return LinkedInJobPosting(
        job_id=job_id,
        title="Product Manager",
        company_name="Acme",
        link=link or f"https://in.linkedin.com/jobs/view/{job_id}",
        location="Bengaluru, India",
        posted_at="2026-08-20",
        applicants_count="50",
        poster_name=poster_name,
        poster_title="Senior HR Executive" if poster_name else None,
        poster_profile_url=f"https://in.linkedin.com/in/{poster_name}" if poster_name else None,
    )


# ---------------------------------------------------------------------------
# extract_emails
# ---------------------------------------------------------------------------


def test_extract_emails_finds_and_dedupes():
    text = "Reach out to jane@acme.com or jane@acme.com, cc hr@acme.io."
    assert extract_emails(text) == ("jane@acme.com", "hr@acme.io")


def test_extract_emails_none_or_empty():
    assert extract_emails(None) == ()
    assert extract_emails("no email here") == ()


# ---------------------------------------------------------------------------
# ApifyLinkedInJobSearch
# ---------------------------------------------------------------------------


def test_apify_search_requires_api_token():
    with pytest.raises(LinkedInJobSearchError, match="token"):
        ApifyLinkedInJobSearch(api_token="")


def test_apify_search_parses_items_into_postings():
    session = FakeSession(
        {
            "Product Manager": FakeResponse(
                json_data=[
                    raw_item("job1", poster_name="Jane Doe", description="email jane@acme.com"),
                    raw_item("job2"),
                ]
            )
        }
    )
    client = ApifyLinkedInJobSearch(api_token="tok", session=session)

    results = client.search("Product Manager", "India", "pastWeek", 10)

    assert len(results) == 2
    assert results[0].company_name == "Acme"
    assert results[0].poster_name == "Jane Doe"
    assert results[0].other_contacts == ("jane@acme.com",)
    assert results[1].poster_name is None

    call = session.calls[0]
    assert call["params"] == {"token": "tok"}
    assert call["json"]["keywords"] == "Product Manager"
    assert call["json"]["location"] == "India"
    assert call["json"]["datePosted"] == "pastWeek"
    assert call["json"]["limitPerSource"] == 10


def test_apify_search_zero_limit_short_circuits_without_a_call():
    session = FakeSession({})
    client = ApifyLinkedInJobSearch(api_token="tok", session=session)

    assert client.search("Product Manager", "India", "pastWeek", 0) == []
    assert session.calls == []


def test_apify_search_non_200_raises():
    session = FakeSession({"Product Manager": FakeResponse(status_code=500)})
    client = ApifyLinkedInJobSearch(api_token="tok", session=session)

    with pytest.raises(LinkedInJobSearchError, match="HTTP 500"):
        client.search("Product Manager", "India", "pastWeek", 10)


def test_apify_search_non_list_response_raises():
    session = FakeSession({"Product Manager": FakeResponse(json_data={"not": "a list"})})
    client = ApifyLinkedInJobSearch(api_token="tok", session=session)

    with pytest.raises(LinkedInJobSearchError, match="Unexpected"):
        client.search("Product Manager", "India", "pastWeek", 10)


# ---------------------------------------------------------------------------
# search_jobs_for_keywords
# ---------------------------------------------------------------------------


def test_search_jobs_for_keywords_caps_combined_total():
    client = FakeJobSearchClient(
        {
            "Product Manager": [posting("pm1"), posting("pm2"), posting("pm3")],
            "APM": [posting("apm1"), posting("apm2")],
        }
    )

    results = search_jobs_for_keywords(
        client, keywords=["Product Manager", "APM"], location="India", date_posted="pastWeek", total_limit=4
    )

    assert [p.job_id for p in results] == ["pm1", "pm2", "pm3", "apm1"]
    # Second keyword's search was asked for only the remaining budget (1), not 2.
    assert client.calls == [
        ("Product Manager", "India", "pastWeek", 4),
        ("APM", "India", "pastWeek", 1),
    ]


def test_search_jobs_for_keywords_dedupes_by_link_across_keywords():
    shared = posting("shared", link="https://in.linkedin.com/jobs/view/shared")
    client = FakeJobSearchClient(
        {
            "Product Manager": [shared],
            "Product Lead": [shared, posting("unique")],
        }
    )

    results = search_jobs_for_keywords(
        client,
        keywords=["Product Manager", "Product Lead"],
        location="India",
        date_posted="pastWeek",
        total_limit=10,
    )

    assert [p.job_id for p in results] == ["shared", "unique"]


def test_search_jobs_for_keywords_stops_once_total_limit_already_hit():
    client = FakeJobSearchClient(
        {
            "Product Manager": [posting("pm1"), posting("pm2")],
            "APM": [posting("apm1")],
        }
    )

    results = search_jobs_for_keywords(
        client, keywords=["Product Manager", "APM"], location="India", date_posted="pastWeek", total_limit=2
    )

    assert [p.job_id for p in results] == ["pm1", "pm2"]
    # APM was never searched — the budget was already exhausted.
    assert client.calls == [("Product Manager", "India", "pastWeek", 2)]


# ---------------------------------------------------------------------------
# write_postings_excel
# ---------------------------------------------------------------------------


def test_write_postings_excel_round_trips(tmp_path: Path):
    rows = [
        LinkedInJobPosting(
            job_id="job1",
            title="Product Manager",
            company_name="Acme",
            link="https://in.linkedin.com/jobs/view/job1",
            location="Bengaluru, India",
            posted_at="2026-08-20",
            applicants_count="50",
            poster_name="Jane Doe",
            poster_title="Senior HR Executive",
            poster_profile_url="https://in.linkedin.com/in/jane-doe",
            other_contacts=("jane@acme.com", "hr@acme.io"),
        ),
        LinkedInJobPosting(
            job_id="job2",
            title="APM",
            company_name="Beta",
            link="https://in.linkedin.com/jobs/view/job2",
            location="Mumbai, India",
            posted_at="2026-08-19",
            applicants_count="12",
            poster_name=None,
            poster_title=None,
            poster_profile_url=None,
        ),
    ]

    output_path = tmp_path / "out" / "jobs.xlsx"
    write_postings_excel(rows, output_path)

    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active

    header = [cell.value for cell in sheet[1]]
    assert header == [
        "Company",
        "Job Title",
        "Job Link",
        "Contact Name",
        "Contact Title",
        "Contact LinkedIn Profile",
        "Other Contact Details",
        "Location",
        "Posted At",
        "Applicants",
    ]

    first_row = [cell.value for cell in sheet[2]]
    assert first_row[0] == "Acme"
    assert first_row[1] == "Product Manager"
    assert first_row[2] == "https://in.linkedin.com/jobs/view/job1"
    assert first_row[3] == "Jane Doe"
    assert first_row[6] == "jane@acme.com; hr@acme.io"

    second_row = [cell.value for cell in sheet[3]]
    assert second_row[3] is None
    assert second_row[6] is None

    assert sheet.cell(row=2, column=3).hyperlink.target == "https://in.linkedin.com/jobs/view/job1"
    assert sheet.cell(row=2, column=6).hyperlink.target == "https://in.linkedin.com/in/jane-doe"


def test_write_postings_excel_empty_list_still_writes_header(tmp_path: Path):
    output_path = tmp_path / "jobs.xlsx"
    write_postings_excel([], output_path)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]][0] == "Company"
