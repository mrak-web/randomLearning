from __future__ import annotations

from pathlib import Path

import openpyxl

from agent.linkedin_export import write_linkedin_excel
from agent.linkedin_jobs import LinkedInJobPosting
from agent.linkedin_posts import HiringPost


def test_write_linkedin_excel_writes_two_sheets(tmp_path: Path):
    job_postings = [
        LinkedInJobPosting(
            job_id="job1",
            title="Product Manager",
            company_name="Acme",
            link="https://in.linkedin.com/jobs/view/job1",
            location="Bengaluru, India",
            posted_at="2026-08-20",
            applicants_count="50",
            poster_name="Jane Doe",
            poster_title="Senior HR Executive",
            poster_profile_url="https://in.linkedin.com/in/jane-doe",
        )
    ]
    hiring_posts = [
        HiringPost(
            post_id="post1",
            author_name="John Smith",
            author_headline="Product Manager @ Beta",
            author_profile_url="https://in.linkedin.com/in/john-smith",
            content="We're hiring a Product Manager, DM for referral.",
            link="https://www.linkedin.com/posts/post1",
            posted_at="2026-08-19",
            company_name="Beta",
        )
    ]

    output_path = tmp_path / "combined.xlsx"
    write_linkedin_excel(job_postings, hiring_posts, output_path)

    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == ["LinkedIn Jobs", "Hiring Posts"]

    jobs_sheet = workbook["LinkedIn Jobs"]
    assert jobs_sheet.cell(row=2, column=1).value == "Acme"

    posts_sheet = workbook["Hiring Posts"]
    assert posts_sheet.cell(row=2, column=1).value == "John Smith"
    assert posts_sheet.cell(row=2, column=4).value == "Beta"


def test_write_linkedin_excel_empty_inputs(tmp_path: Path):
    output_path = tmp_path / "empty.xlsx"
    write_linkedin_excel([], [], output_path)

    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == ["LinkedIn Jobs", "Hiring Posts"]
    assert workbook["LinkedIn Jobs"].max_row == 1
    assert workbook["Hiring Posts"].max_row == 1
