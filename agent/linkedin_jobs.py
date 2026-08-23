"""LinkedIn job-posting lookup — a standalone tool, not part of the cold-email
pipeline. See PROJECT.md's "LinkedIn Job Lookup" section.

Searches LinkedIn's *public, logged-out* jobs search page via the Apify Actor
curious_coder/linkedin-jobs-scraper. This never authenticates as Arjun's own
LinkedIn account, so unlike the scraping avoided elsewhere in this project (see
PROJECT.md §7's LinkedIn risk note, which is about *account* scraping), there is no
ToS/ban risk to his profile — the Actor's own infrastructure does the request, and
the search endpoint it hits is the same one any signed-out visitor can reach.

LinkedInJobSearchClient is a thin interface (same swap pattern as EmailFinder/
CompanySource/EmailSender elsewhere in agent/) so orchestration logic here is
testable against a fake, with no live Apify credits spent in tests.
"""

from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import requests
from openpyxl import Workbook

APIFY_RUN_SYNC_URL_TEMPLATE = "https://api.apify.com/v2/acts/{actor_slug}/run-sync-get-dataset-items"

# Loose but practical email matcher — good enough to catch an address someone typed
# into a job description's free text (e.g. "reach out to jane@acme.com"), which is
# the only place a job posting is likely to expose one directly.
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

EXCEL_HEADERS = [
    "Company",
    "Job Title",
    "Match Score",
    "Match Notes",
    "Job Link",
    "Contact Name",
    "Contact Title",
    "Contact LinkedIn Profile",
    "Other Contact Details",
    "Location",
    "Posted At",
    "Applicants",
]


class LinkedInJobSearchError(Exception):
    """Raised when the Apify Actor call fails or returns something unusable."""


@dataclass(frozen=True)
class LinkedInJobPosting:
    job_id: str | None
    title: str | None
    company_name: str | None
    link: str | None
    location: str | None
    posted_at: str | None
    applicants_count: str | None
    poster_name: str | None
    poster_title: str | None
    poster_profile_url: str | None
    other_contacts: tuple[str, ...] = ()
    description_text: str | None = None
    # Filled in by agent.resume_match.attach_match_scores after search — kept on the
    # same record (rather than a parallel lookup) so the row that flows into the
    # Excel writer is already fully enriched, same pattern as other_contacts above.
    match_score: int | None = None
    match_reasons: tuple[str, ...] = ()


def extract_emails(text: str | None) -> tuple[str, ...]:
    """Pulls email addresses out of free text, deduped, first-seen order preserved."""
    if not text:
        return ()
    seen: list[str] = []
    for match in EMAIL_PATTERN.findall(text):
        if match not in seen:
            seen.append(match)
    return tuple(seen)


def _to_posting(item: dict) -> LinkedInJobPosting:
    return LinkedInJobPosting(
        job_id=item.get("id"),
        title=item.get("title"),
        company_name=item.get("companyName"),
        link=item.get("link"),
        location=item.get("location"),
        posted_at=item.get("postedAt"),
        applicants_count=item.get("applicantsCount"),
        poster_name=item.get("jobPosterName"),
        poster_title=item.get("jobPosterTitle"),
        poster_profile_url=item.get("jobPosterProfileUrl"),
        other_contacts=extract_emails(item.get("descriptionText")),
        description_text=item.get("descriptionText"),
    )


class LinkedInJobSearchClient(ABC):
    """A pluggable job-search source. Apify's curious_coder Actor is the only one today."""

    @abstractmethod
    def search(
        self, keywords: str, location: str, date_posted: str, limit: int
    ) -> list[LinkedInJobPosting]:
        """Runs one search for a single free-text keyword string, capped at `limit` results."""


class ApifyLinkedInJobSearch(LinkedInJobSearchClient):
    """Calls curious_coder/linkedin-jobs-scraper via Apify's run-sync REST endpoint.

    Uses the synchronous run-sync-get-dataset-items endpoint (one HTTP call per
    keyword, blocks until the run finishes) rather than the async run+poll API — the
    result sets here are small enough (<=~100 items) that this stays well within
    Apify's ~5 minute sync-endpoint window.
    """

    actor_slug = "curious_coder~linkedin-jobs-scraper"

    def __init__(
        self,
        api_token: str,
        session: requests.Session | None = None,
        timeout: float = 280.0,
    ):
        if not api_token:
            raise LinkedInJobSearchError("Apify API token is required")
        self.api_token = api_token
        self.session = session or requests.Session()
        self.timeout = timeout

    def search(
        self, keywords: str, location: str, date_posted: str, limit: int
    ) -> list[LinkedInJobPosting]:
        if limit <= 0:
            return []

        url = APIFY_RUN_SYNC_URL_TEMPLATE.format(actor_slug=self.actor_slug)
        payload = {
            "keywords": keywords,
            "location": location,
            "datePosted": date_posted,
            "limitPerSource": limit,
            # Company detail scraping is a separate per-job request in the Actor —
            # skipped here since it's not one of the requested Excel columns.
            "scrapeCompany": False,
        }
        try:
            response = self.session.post(
                url, params={"token": self.api_token}, json=payload, timeout=self.timeout
            )
        except requests.RequestException as exc:
            raise LinkedInJobSearchError(
                f"Apify request failed for keywords={keywords!r}: {exc}"
            ) from exc

        # Apify's run-sync-get-dataset-items endpoint replies 201 (Created) on a
        # normal successful run, not 200 — it's creating a run resource, not just
        # returning data. Any 2xx is success; only treat 3xx/4xx/5xx as failure.
        if not (200 <= response.status_code < 300):
            raise LinkedInJobSearchError(
                f"Apify Actor call failed for keywords={keywords!r}: HTTP {response.status_code}"
            )

        try:
            items = response.json()
        except ValueError as exc:
            raise LinkedInJobSearchError(
                f"Apify returned a non-JSON response for keywords={keywords!r}"
            ) from exc

        if not isinstance(items, list):
            raise LinkedInJobSearchError(
                f"Unexpected Apify response shape for keywords={keywords!r}"
            )

        return [_to_posting(item) for item in items]


def search_jobs_for_keywords(
    client: LinkedInJobSearchClient,
    keywords: Iterable[str],
    location: str,
    date_posted: str,
    total_limit: int,
) -> list[LinkedInJobPosting]:
    """Runs one search per keyword, deduped by job link, capped at total_limit overall.

    The cap applies to the *combined* result count across all keywords, not per
    keyword — this keeps Apify spend predictable (proportional to total_limit)
    regardless of how many keywords settings.yaml lists.
    """
    collected: list[LinkedInJobPosting] = []
    seen_keys: set[str] = set()

    for keyword in keywords:
        remaining = total_limit - len(collected)
        if remaining <= 0:
            break

        for posting in client.search(keyword, location, date_posted, remaining):
            key = posting.link or posting.job_id
            if key and key in seen_keys:
                continue
            if key:
                seen_keys.add(key)
            collected.append(posting)
            if len(collected) >= total_limit:
                break

    return collected


def autosize_columns(sheet) -> None:
    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 60)


def populate_jobs_sheet(sheet, postings: list[LinkedInJobPosting]) -> None:
    """Fills an already-created worksheet with the job-listing columns and autosizes it.

    Factored out of write_postings_excel so agent/linkedin_export.py can put this
    sheet alongside the separate "Hiring Posts" sheet in one combined workbook.
    """
    sheet.append(EXCEL_HEADERS)

    for posting in postings:
        row = [
            posting.company_name,
            posting.title,
            posting.match_score,
            "; ".join(posting.match_reasons) or None,
            posting.link,
            posting.poster_name,
            posting.poster_title,
            posting.poster_profile_url,
            "; ".join(posting.other_contacts) or None,
            posting.location,
            posting.posted_at,
            posting.applicants_count,
        ]
        sheet.append(row)
        written_row = sheet.max_row
        if posting.link:
            sheet.cell(row=written_row, column=5).hyperlink = posting.link
        if posting.poster_profile_url:
            sheet.cell(row=written_row, column=8).hyperlink = posting.poster_profile_url

    autosize_columns(sheet)


def write_postings_excel(postings: list[LinkedInJobPosting], output_path: Path) -> None:
    """Writes postings to a single-sheet .xlsx with the columns Arjun asked for."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LinkedIn Jobs"
    populate_jobs_sheet(sheet, postings)

    workbook.save(output_path)
